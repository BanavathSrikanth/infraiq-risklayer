import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.domain.schemas.column_mapping import ColumnProfile


def profile_file(filename: str, content: bytes, sample_size: int = 100) -> list[ColumnProfile]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))[:sample_size]
    elif suffix == ".json":
        payload = json.loads(content)
        rows = payload if isinstance(payload, list) else payload.get("rows", [payload])
    elif suffix in {".xlsx", ".xls"}:
        rows = _excel_rows(content, sample_size, suffix)
    elif suffix == ".parquet":
        rows = _parquet_rows(content, sample_size)
    else:
        raise ValueError("Supported file types are CSV, Excel, JSON, and Parquet")

    if not rows:
        return []
    names = list(rows[0].keys())
    profiles = []
    for name in names:
        values = [row.get(name) for row in rows]
        non_null = [value for value in values if value not in (None, "")]
        profiles.append(
            ColumnProfile(
                name=name,
                inferred_type=_infer_type(non_null),
                nullable=len(non_null) != len(values),
                sample_values=[str(value)[:200] for value in non_null[:5]],
                null_fraction=1 - (len(non_null) / len(values)),
            )
        )
    return profiles


def _excel_rows(content: bytes, sample_size: int, suffix: str) -> list[dict[str, Any]]:
    if suffix == ".xls":
        import pandas as pd

        return pd.read_excel(io.BytesIO(content), engine="xlrd").head(sample_size).to_dict(orient="records")

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    return [dict(zip(headers, row)) for row in list(rows)[:sample_size]]


def _parquet_rows(content: bytes, sample_size: int) -> list[dict[str, Any]]:
    import pandas as pd

    return pd.read_parquet(io.BytesIO(content)).head(sample_size).to_dict(orient="records")


def _infer_type(values: list[Any]) -> str:
    if not values:
        return "empty"
    if all(isinstance(value, bool) for value in values):
        return "boolean"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in values):
        return "integer"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
        return "number"
    if all(isinstance(value, (date, datetime)) for value in values):
        return "date"
    if all(_is_date_string(value) for value in values):
        return "date"
    if all(_is_integer_string(value) for value in values):
        return "integer"
    if all(_is_number_string(value) for value in values):
        return "number"
    return "string"


def _is_date_string(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    try:
        datetime.fromisoformat(value.replace("Z", "+00:00"))
        return True
    except ValueError:
        return False


def _is_integer_string(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    try:
        int(value)
        return True
    except ValueError:
        return False


def _is_number_string(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    try:
        float(value)
        return True
    except ValueError:
        return False